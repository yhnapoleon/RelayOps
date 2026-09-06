"""Control-M job-sheet import — the second half of a Job's identity.

The handover checklist gives the MMP model / scenarios; the *real* Control-M
job names + schedules live in a separate config sheet (a Confluence table or
an Excel/CSV the EJS team maintains), e.g.:

    Job Name                                        | Schedule (free text)              | PARM1 (Project_Name)
    PKG_CML_MATERIAL_CLASSIFIER_NS_SG_RUN_W_EDSP    | Run weekly at SGT 2:55pm Monday   | material-classifier

This module turns that sheet into ``control_m_job_name`` + ``schedule_cron``
on the draft's jobs:

  1. parse the upload to plain text (xlsx via openpyxl, else csv/tsv/txt/html);
  2. LLM-extract only the *production scheduled* jobs (adhoc / rerun /
     regression / standby rows are dropped), with the cron derived from the
     free-text schedule;
  3. deterministically merge into the draft — fill an existing job that lacks a
     Control-M name when it clearly corresponds, else add the job. Nothing the
     reviewer already set is overwritten.

Used as a per-draft enrichment action (upload onto an existing draft), not a
new-project ingest.
"""

from __future__ import annotations

import csv
import difflib
import io
import re
from typing import List, Optional

from pydantic import BaseModel, Field

from core.agent.schemas import JobDraft, OnboardingDraftPayload, ProductDraft
from core.logging import get_logger

logger = get_logger(__name__)

_MERGE_FLOOR = 0.7   # below this, don't fold a sheet job into an existing one


class UnsupportedSheetError(ValueError):
    """File extension we can't turn into text — caller asks for csv/paste."""


# ── LLM output schema ─────────────────────────────────────────────────


class ControlMJob(BaseModel):
    control_m_job_name: str = ""
    cml_project_name: str = ""
    schedule_cron: str = ""
    description: str = ""


class ControlMSheet(BaseModel):
    jobs: List[ControlMJob] = Field(default_factory=list)


# ── parsing the upload to text ────────────────────────────────────────


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "gbk"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def parse_sheet_to_text(filename: str, data: bytes) -> str:
    """Flatten the uploaded sheet to ``cell | cell`` rows for the LLM.

    .xlsx → openpyxl (agent extra). .csv/.tsv/.txt → decoded and re-joined.
    .html/.htm → the shared HTML flattener. Other extensions raise.
    """
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook  # agent extra
        except ImportError as exc:
            raise UnsupportedSheetError(
                "Parsing .xlsx requires openpyxl (agent extra); please install it "
                "and retry, or export to CSV / paste the text"
            ) from exc
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise UnsupportedSheetError(f"Excel parse failed: {exc}") from exc
        lines: List[str] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c).strip() for c in row]
                if any(cells):
                    lines.append(" | ".join(cells))
        return "\n".join(lines)[:200_000]
    if name.endswith((".html", ".htm")):
        from core.agent.ingest import html_to_text

        return html_to_text(_decode(data))
    if name.endswith((".csv", ".tsv", ".txt")) or not name:
        text = _decode(data)
        # Normalize CSV/TSV rows to "cell | cell" so column semantics survive.
        delim = "\t" if (name.endswith(".tsv") or "\t" in text) else ","
        out: List[str] = []
        for row in csv.reader(io.StringIO(text), delimiter=delim):
            cells = [c.strip() for c in row]
            if any(cells):
                out.append(" | ".join(cells))
        return ("\n".join(out) or text)[:200_000]
    raise UnsupportedSheetError(
        f"Unsupported file format: {filename!r} (supported: xlsx / csv / tsv / txt / html, or paste directly)"
    )


# ── LLM extraction ────────────────────────────────────────────────────


def extract_controlm_jobs(sheet_text: str) -> List[ControlMJob]:
    """One structured-output call → the sheet's production scheduled jobs.
    Raises on failure (caller marks the draft failed)."""
    from core.agent.llm import get_chat_model
    from core.agent.prompts import CONTROLM_EXTRACT_SYSTEM
    from core.agent.structured import invoke_structured

    chat = get_chat_model(temperature=0)
    result = invoke_structured(chat, ControlMSheet, [
        ("system", CONTROLM_EXTRACT_SYSTEM),
        ("human", sheet_text[:120_000]),
    ])
    # Keep only rows with a real job name; the LLM is told to skip adhoc/rerun.
    return [j for j in result.jobs if j.control_m_job_name.strip()]


# ── deterministic merge into the draft ────────────────────────────────


def _norm(s: str) -> str:
    return re.sub(r"[\s_\-./]+", "", (s or "").lower())


def _score(a: str, b: str) -> float:
    na, nb = _norm(a), _norm(b)
    if not na or not nb:
        return 0.0
    if na == nb:
        return 1.0
    base = difflib.SequenceMatcher(None, na, nb).ratio()
    if na in nb or nb in na:
        base = max(base, 0.85)
    return base


def _ensure_product(payload: OnboardingDraftPayload) -> ProductDraft:
    if payload.products:
        return payload.products[0]
    name = payload.project.name.strip() or payload.project.cml_project_name.strip() or "Default Product"
    product = ProductDraft(name=name)
    payload.products.append(product)
    return product


def _all_jobs(payload: OnboardingDraftPayload):
    for prod in payload.products:
        for job in prod.jobs:
            yield job


def _llm_match_controlm_to_jobs(
    source_text: str, cm_jobs: List[ControlMJob], draft_jobs: List
) -> List[int]:
    """LLM maps each Control-M job to the draft job index it fills (or -1 = add
    new), using the source document for the region ↔ MMP-model cross-reference
    that pure string similarity can't see. Returns one index per ``cm_jobs``
    entry (same order); all -1 on any failure (caller falls back to similarity).
    """
    if not cm_jobs or not draft_jobs or not source_text.strip():
        return [-1 for _ in cm_jobs]
    try:
        from core.agent.llm import get_chat_model
        from core.agent.prompts import CONTROLM_MATCH_SYSTEM
        from core.agent.structured import invoke_structured

        class _Assign(BaseModel):
            assignments: List[int] = Field(default_factory=list)

        draft_desc = "\n".join(
            f"{i}. mmp_model={j.mmp_model_id or '-'} | cml_job={j.cml_job_name or '-'} "
            f"| description={(j.description or j.source_quote or '')[:160]!r}"
            for i, j in enumerate(draft_jobs)
        )
        cm_desc = "\n".join(
            f"{i}. name={c.control_m_job_name!r} | description={c.description[:120]!r}"
            for i, c in enumerate(cm_jobs)
        )
        chat = get_chat_model(temperature=0)
        result = invoke_structured(chat, _Assign, [
            ("system", CONTROLM_MATCH_SYSTEM),
            ("human", f"Document snippet:\n{source_text[:6000]}\n\n"
                      f"Draft jobs to fill:\n{draft_desc}\n\n"
                      f"Control-M jobs to place:\n{cm_desc}"),
        ])
        picks = list(result.assignments) + [-1] * len(cm_jobs)
        out: List[int] = []
        for idx in picks[:len(cm_jobs)]:
            out.append(idx if isinstance(idx, int) and 0 <= idx < len(draft_jobs) else -1)
        return out
    except Exception:  # noqa: BLE001 — matching must never fail the merge
        logger.opt(exception=True).warning("onboarding: LLM Control-M match failed")
        return [-1 for _ in cm_jobs]


def merge_controlm_into_draft(payload: OnboardingDraftPayload,
                              cm_jobs: List[ControlMJob],
                              source_text: str = "") -> List[str]:
    """Fill / add jobs from the sheet. Returns human-readable notes (warnings).

    For each sheet job: if its Control-M name is already on a draft job, only
    backfill an empty schedule; else fold it into the corresponding draft job
    that still lacks a Control-M name — first via an LLM that reads the source
    document's region ↔ MMP-model mapping, then a string-similarity fallback;
    else add it as a new job. Reviewer values are never overwritten.
    """
    notes: List[str] = []
    project_cml = payload.project.cml_project_name.strip()

    # LLM pass: map the not-yet-present sheet jobs to unbound draft jobs by the
    # document's cross-reference (e.g. MMP model 344 ↔ "NS GE"), which string
    # similarity alone misses. Falls back to per-job similarity below.
    pending = [cm for cm in cm_jobs if cm.control_m_job_name.strip()
               and not next((j for j in _all_jobs(payload)
                             if _norm(j.control_m_job_name) == _norm(cm.control_m_job_name)), None)]
    unbound = [j for j in _all_jobs(payload) if not j.control_m_job_name.strip()]
    llm_pick = {id(cm): i for cm, i in zip(
        pending, _llm_match_controlm_to_jobs(source_text, pending, unbound))}
    llm_taken: set = set()

    for cm in cm_jobs:
        name = cm.control_m_job_name.strip()
        if not name:
            continue

        # Already present? (idempotent re-upload) — just backfill schedule.
        existing = next((j for j in _all_jobs(payload)
                         if _norm(j.control_m_job_name) == _norm(name)), None)
        if existing is not None:
            if cm.schedule_cron.strip() and not existing.schedule_cron.strip():
                existing.schedule_cron = cm.schedule_cron.strip()
                notes.append(f"Backfilled the schedule cron for existing job 「{name}」: {cm.schedule_cron.strip()}")
            continue

        # Prefer the LLM's document-aware assignment (each draft job once).
        best, best_score = None, 0.0
        pick = llm_pick.get(id(cm), -1)
        if 0 <= pick < len(unbound) and pick not in llm_taken \
                and not unbound[pick].control_m_job_name.strip():
            best, best_score = unbound[pick], 1.0
            llm_taken.add(pick)
        else:
            # Fallback: fold into the best similar unbound job (same/empty CML
            # project + similar description or model), above a conservative floor.
            for job in _all_jobs(payload):
                if job.control_m_job_name.strip():
                    continue
                jproj = job.cml_project_name.strip() or project_cml
                if cm.cml_project_name.strip() and jproj and \
                        _norm(jproj) != _norm(cm.cml_project_name):
                    continue  # bound to a different CML project
                score = max(_score(cm.description, job.description),
                            _score(cm.description, job.mmp_model_id),
                            _score(cm.description, job.cml_job_name))
                if score > best_score:
                    best, best_score = job, score

        if best is not None and best_score >= _MERGE_FLOOR:
            best.control_m_job_name = name
            if cm.schedule_cron.strip() and not best.schedule_cron.strip():
                best.schedule_cron = cm.schedule_cron.strip()
            if cm.cml_project_name.strip() and not best.cml_project_name.strip() \
                    and _norm(cm.cml_project_name) != _norm(project_cml):
                best.cml_project_name = cm.cml_project_name.strip()
            notes.append(
                f"Merged Control-M job 「{name}」 ({cm.description or 'no description'}) "
                f"into the matching job and backfilled schedule "
                f"{cm.schedule_cron or '(not inferred)'} — please verify"
            )
        else:
            override = (cm.cml_project_name.strip()
                        if _norm(cm.cml_project_name) != _norm(project_cml) else "")
            _ensure_product(payload).jobs.append(JobDraft(
                control_m_job_name=name,
                schedule_cron=cm.schedule_cron.strip(),
                cml_project_name=override,
                description=cm.description.strip(),
                source_quote=f"Control-M sheet: {name}",
            ))
            notes.append(
                f"Created job 「{name}」 from the Control-M sheet "
                f"({cm.description or 'no description'}), schedule "
                f"{cm.schedule_cron or '(not inferred — please add)'} — please verify"
            )

    if cm_jobs:
        notes.append(
            f"Imported {len(cm_jobs)} production jobs from the Control-M sheet "
            "(adhoc/rerun/regression skipped); every inferred cron must be "
            "verified against the actual Control-M schedule"
        )
    return notes
